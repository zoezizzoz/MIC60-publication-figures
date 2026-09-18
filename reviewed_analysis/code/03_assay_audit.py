"""Recalculate accessible assay summaries without treating subsamples as replicates.

Run with the bundled Python runtime documented in ../README.md.
All input workbooks are read only; outputs go to ../results.
"""
from pathlib import Path
import itertools
import json
import numpy as np
import pandas as pd
import subprocess
import os
from types import SimpleNamespace

def r_eval(code):
    r = subprocess.run([os.environ.get("RSCRIPT", "Rscript"), "-e", "suppressPackageStartupMessages(library(jsonlite)); "+code], capture_output=True, text=True, check=True)
    return json.loads(r.stdout)

def rv(x):
    return "c("+",".join(format(float(v),".17g") for v in x)+")"

class RStats:
    @staticmethod
    def f_oneway(*groups):
        ans=r_eval("x<-"+rv(np.concatenate(groups))+";g<-factor(rep(1:"+str(len(groups))+",c("+",".join(str(len(g)) for g in groups)+")));a<-summary(aov(x~g))[[1]];cat(toJSON(list(statistic=a[1,4],pvalue=a[1,5]),auto_unbox=TRUE,digits=16))")
        return SimpleNamespace(**ans)
    @staticmethod
    def tukey_hsd(*groups):
        ans=r_eval("x<-"+rv(np.concatenate(groups))+";g<-factor(rep(1:"+str(len(groups))+",c("+",".join(str(len(g)) for g in groups)+")));a<-TukeyHSD(aov(x~g))$g;cat(toJSON(data.frame(pair=rownames(a),p=a[,4]),digits=16))")
        p=np.ones((len(groups),len(groups)))
        for row in ans:
            i,j=[int(v)-1 for v in row["pair"].split("-")];p[i,j]=p[j,i]=row["p"]
        return SimpleNamespace(pvalue=p)
    @staticmethod
    def mannwhitneyu(x,y,**kwargs):
        exact = "TRUE" if kwargs.get("method")=="exact" else "FALSE"
        a=r_eval("a<-wilcox.test("+rv(x)+","+rv(y)+",exact="+exact+",correct=TRUE);cat(toJSON(list(pvalue=a$p.value),auto_unbox=TRUE,digits=16))")
        return SimpleNamespace(**a)
    @staticmethod
    def rankdata(x):
        return pd.Series(x).rank().to_numpy()
    @staticmethod
    def fisher_exact(table,**kwargs):
        a=r_eval("a<-fisher.test(matrix("+rv(np.array(table).flatten())+",2,byrow=TRUE),alternative=\"greater\");cat(toJSON(list(statistic=unname(a$estimate),pvalue=a$p.value),auto_unbox=TRUE,digits=16,na=\"string\"))")
        a["statistic"]=float(a["statistic"])
        return SimpleNamespace(**a)

stats=RStats()
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA, OUT = ROOT / 'data', ROOT / 'results'
OUT.mkdir(exist_ok=True)
results = {}

def bh(p):
    p = np.asarray(p, dtype=float)
    order = np.argsort(p)
    adj = np.minimum.accumulate((p[order] * len(p) / np.arange(1, len(p)+1))[::-1])[::-1]
    out = np.empty(len(p)); out[order] = np.minimum(adj, 1)
    return out

# Fig 4A: replicate identities are not supplied in this four-column workbook.
w = load_workbook(DATA/'TIMELESS_DAPI.xlsx', data_only=True, read_only=True)
rows = list(w.active.values)
groups = {h: np.array([r[i] for r in rows[1:] if isinstance(r[i], (int,float))])
          for i,h in enumerate(rows[0])}
anova = stats.f_oneway(*groups.values())
tukey = stats.tukey_hsd(*groups.values())
summary = pd.DataFrame([{'group': k, 'observations': len(v), 'mean': v.mean(),
                        'median': np.median(v), 'sd': v.std(ddof=1)} for k,v in groups.items()])
summary.to_csv(OUT/'TIMELESS_DAPI_descriptive_summary.csv', index=False)
comparisons=[]
for i,j in itertools.combinations(range(4),2):
    comparisons.append({'group1': list(groups)[i], 'group2': list(groups)[j],
                        'Tukey_p_assuming_independent_observations': float(tukey.pvalue[i,j])})
pd.DataFrame(comparisons).to_csv(OUT/'TIMELESS_DAPI_conditional_test_audit.csv',index=False)
results['TIMELESS_DAPI']={'F':float(anova.statistic),'df_between':3,
    'df_within':sum(map(len,groups.values()))-4, 'P':float(anova.pvalue),
    'limitation':'Conditional numerical reproduction only. Biological replicate IDs are absent; these tests do not establish replicate-level inference.'}

# Fig 4B: reproduce the normalized values from raw absorbance and blanks.
w=load_workbook(DATA/'MTT_source.xlsx',data_only=True,read_only=True)
s=w['7.28.26 raw']; records=[]
blank=np.mean([s.cell(35,c).value for c in (2,3,4)])
# Raw plate columns: A-F rows; col 1 is the row letter.
for dose,cols in [(0,(4,5)),(5,(12,13)),(10,(10,11)),(20,(8,9)),(40,(6,7))]:
    for geno,col in zip(('WT','CS'),cols):
        for well,row in enumerate(range(29,35),1):
            records.append({'experiment':'7.28.26','dose_mM':dose,'genotype':geno,'well':well,
                            'raw_absorbance':s.cell(row,col).value,'blank':blank})
s=w['7.15.26 raw'];blank=np.mean([s.cell(6,c).value for c in (3,4)])
for dose,cols in [(0,(4,5)),(20,(6,7)),(40,(8,9))]:
    for geno,col in zip(('WT','CS'),cols):
        for well,row in enumerate(range(3,6),1):
            records.append({'experiment':'7.15.26','dose_mM':dose,'genotype':geno,'well':well,
                            'raw_absorbance':s.cell(row,col).value,'blank':blank})
d=pd.DataFrame(records);d['corrected_absorbance']=d.raw_absorbance-d.blank
controls=d[d.dose_mM==0].groupby(['experiment','genotype']).corrected_absorbance.mean()
d['normalized_pct']=[100*r.corrected_absorbance/controls.loc[(r.experiment,r.genotype)] for r in d.itertuples()]
saved=load_workbook(DATA/'MTT_plotted.xlsx',data_only=True,read_only=True)
p=pd.DataFrame(list(saved['Individual datapoints'].values)[4:],columns=['experiment','dose_mM','genotype','label','well','saved_pct']).dropna(subset=['experiment'])
check=d.merge(p,on=['experiment','dose_mM','genotype','well'],validate='one_to_one')
assert len(check)==78
maxdiff=float(np.max(np.abs(check.normalized_pct-check.saved_pct)))
assert maxdiff<1e-10, f'MTT does not reconcile: {maxdiff}'
d.to_csv(OUT/'MTT_reconstructed_wells.csv',index=False)
means=d.groupby(['experiment','dose_mM','genotype'],as_index=False).agg(mean_pct=('normalized_pct','mean'),technical_n=('normalized_pct','size'))
means.to_csv(OUT/'MTT_experiment_means.csv',index=False)
means.groupby(['dose_mM','genotype']).agg(mean_pct=('mean_pct','mean'),sd_between_experiments=('mean_pct','std'),independent_experiments=('mean_pct','size'),technical_wells=('technical_n','sum')).to_csv(OUT/'MTT_descriptive_summary.csv')
results['MTT']={'raw_to_plotted_max_absolute_error':maxdiff,'technical_wells':len(d),
 'limitation':'One experiment at 5/10 mM; two at 0/20/40 mM. No well-level significance annotations. July 15 source control is labelled DMSO; confirm vehicle before interpreting pooled experiments.'}

# TEM: field-level tests can be reproduced but fly IDs are absent.
d=pd.read_csv(DATA/'TEM_fields.csv');tem=[]
for metric,g in d.groupby('Metric'):
    x=g.loc[g.Condition=='WR','Image_Mean'].to_numpy();y=g.loc[g.Condition=='CS','Image_Mean'].to_numpy()
    if not len(x):x=g.loc[g.Condition=='WT','Image_Mean'].to_numpy()
    method='exact' if len(np.unique(np.r_[x,y]))==len(x)+len(y) else 'asymptotic'
    r=stats.mannwhitneyu(x,y,method=method,use_continuity=True)
    tem.append({'metric':metric,'WT_fields':len(x),'CS_fields':len(y),'WT_mean':x.mean(),'CS_mean':y.mean(),
                'P_field_level_exploratory':r.pvalue})
pd.DataFrame(tem).to_csv(OUT/'TEM_field_level_audit.csv',index=False)

# Geotaxis: ties invalidate base R's default untied exact Wilcoxon distribution.
d=pd.read_csv(DATA/'locomotion.csv');x=d[d.group=='WT'].performance_index.to_numpy();y=d[d.group=='CS'].performance_index.to_numpy()
pooled=np.r_[x,y];ranks=stats.rankdata(pooled);observed=ranks[:len(x)].sum();center=len(x)*(len(pooled)+1)/2
perms=[ranks[list(ix)].sum() for ix in itertools.combinations(range(len(pooled)),len(x))]
p_exact=np.mean(np.abs(np.array(perms)-center)>=abs(observed-center)-1e-12)
results['locomotion']={'WT_n':len(x),'CS_n':len(y),'exact_two_sided_rank_permutation_P':float(p_exact),
 'asymptotic_tie_corrected_P':float(stats.mannwhitneyu(x,y,method='asymptotic').pvalue),
 'note':'Exact permutation enumerates all 70 allocations and handles ties explicitly; inference assumes independent populations.'}

# Stress sets: full precision Fisher P values; BH across the six prespecified displays.
d=pd.read_csv(DATA/'stress_genes.csv');universe=d.drop_duplicates('gene').dropna(subset=['padj_F']);up=set(universe.loc[(universe.padj_F<.05)&(universe.lfc_F>=.58),'gene']);U=set(universe.gene);sets=[]
for pathway,g in d.groupby('pathway',sort=False):
    genes=set(g.gene)&U;outside=U-genes
    a,b=len(genes&up),len(genes-up);c,e=len(outside&up),len(outside-up)
    r=stats.fisher_exact([[a,b],[c,e]],alternative='greater')
    sets.append({'pathway':pathway,'testable_genes':len(genes),'upregulated':a,'odds_ratio':r.statistic,'P':r.pvalue})
sets=pd.DataFrame(sets);sets['BH_adjusted_P']=bh(sets.P);sets.to_csv(OUT/'stress_enrichment_audit.csv',index=False)
results['stress']={'unique_testable_genes':len(U),'upregulated_genes':sorted(up),
 'limitation':'Enrichment relative to other curated genes; nonsignificance does not show absence of pathway activation. ER-UPR sets are not a dedicated mitochondrial UPR assay.'}

for sex in ('female','male'):
    d=pd.read_csv(DATA/f'rnaseq_{sex}.csv');test=d.dropna(subset=['padj']);up=test[(test.padj<.05)&(test.log2FoldChange>=.58)];down=test[(test.padj<.05)&(test.log2FoldChange<=-.58)]
    results[sex+'_DE']={'testable_genes':len(test),'up':len(up),'down':len(down),
     'selected_genes':d[d.gene.isin(['timeout','DNAlig3','ImpL2'])][['gene','log2FoldChange','padj']].to_dict('records')}
(OUT/'assay_audit.json').write_text(json.dumps(results,indent=2,allow_nan=False)+'\n')
print(json.dumps(results,indent=2))
